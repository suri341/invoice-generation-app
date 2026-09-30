import io
from datetime import date

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def xlsx_response(title: str, headers: list, rows: list) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, 1):
        width = max([len(str(header))] + [len(str(row[index - 1] or "")) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 60)
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = f"{title.lower()}_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
